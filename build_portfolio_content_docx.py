import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(r"E:\onedrive\claude\tomob")
SOURCE = ROOT / "project" / "newseason" / "전성은_포트폴리오_전체원고_Claude전달용.md"
OUTPUT = ROOT / "output" / "documents" / "전성은_포트폴리오_전체원고_Claude전달용.docx"

FONT = "맑은 고딕"
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(93, 102, 115)
ACCENT = RGBColor(158, 67, 55)
NAVY = RGBColor(28, 60, 88)


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    set_font(run, 8, True, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_rich_text(paragraph, text, size=10.5, color=INK):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, 9, False, MUTED)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size, True, color)
        else:
            run = paragraph.add_run(part)
            set_font(run, size, False, color)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(18)
    sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(20)
    sec.right_margin = Mm(20)
    sec.header_distance = Mm(8)
    sec.footer_distance = Mm(9)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 18, NAVY, 18, 9),
        ("Heading 2", 14, NAVY, 14, 7),
        ("Heading 3", 11.5, ACCENT, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.text = "전성은 포트폴리오 전체 원고  ·  Claude 전달용"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_font(run, 8, True, MUTED)

    add_page_field(sec.footer.paragraphs[0])


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)

    first_title = True
    seen_major = False
    for raw in lines:
        line = raw.strip()
        if not line or line == "---":
            continue

        if line.startswith("# "):
            text = line[2:].strip()
            if first_title:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(54)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run(text)
                set_font(r, 25, True, NAVY)
                first_title = False
                continue
            needs_break = (
                text.startswith("PAGE 01.")
                or text.startswith("PROJECT ")
                or text.startswith("부록 ")
            )
            if needs_break and seen_major:
                doc.add_page_break()
            if needs_break:
                seen_major = True
            p = doc.add_paragraph(style="Heading 1")
            add_rich_text(p, text, 18, NAVY)
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_rich_text(p, line[3:].strip(), 14, NAVY)
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_rich_text(p, line[4:].strip(), 11.5, ACCENT)
            continue

        if re.match(r"^-\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Mm(6.4)
            p.paragraph_format.first_line_indent = Mm(-3.2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            add_rich_text(p, re.sub(r"^-\s+", "", line))
            continue

        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Mm(7)
            p.paragraph_format.first_line_indent = Mm(-3.5)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            add_rich_text(p, re.sub(r"^\d+\.\s+", "", line))
            continue

        p = doc.add_paragraph()
        if line.startswith("“") or line.startswith('"'):
            p.paragraph_format.left_indent = Mm(7)
            p.paragraph_format.right_indent = Mm(7)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            add_rich_text(p, line, 10.5, MUTED)
            for run in p.runs:
                run.italic = True
        else:
            add_rich_text(p, line)

    props = doc.core_properties
    props.title = "전성은 포트폴리오 전체 원고"
    props.subject = "Claude 전달용 콘텐츠 정본"
    props.author = "전성은"
    props.keywords = "포트폴리오, 서비스기획, 프로덕트기획, UXUI, AX"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
